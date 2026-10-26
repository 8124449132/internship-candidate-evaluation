
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify
from openpyxl import load_workbook
from pathlib import Path
from urllib.parse import urlparse
import os
import re
import base64
import requests
from openai import OpenAI

app = Flask(__name__)
app.secret_key = os.environ.get(
    "FLASK_SECRET_KEY", "change-this-for-local-demo"
)

BASE_DIR = Path(__file__).resolve().parent
EXCEL_PATH = BASE_DIR / "Internship_Candidates_100.xlsx"
SHEET_NAME = "Candidates"


def get_sheet():
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            "Internship_Candidates_100.xlsx not found. "
            "Put the Excel file in the same folder as app.py."
        )

    wb = load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    return wb, ws


def get_candidates():
    wb, ws = get_sheet()
    headers = [cell.value for cell in ws[1]]
    candidates = []

    for row_num in range(2, ws.max_row + 1):
        values = [
            ws.cell(row_num, col).value
            for col in range(1, len(headers) + 1)
        ]

        if not any(value is not None for value in values):
            continue

        item = dict(zip(headers, values))
        item["_row"] = row_num

        for key, value in list(item.items()):
            if hasattr(value, "strftime"):
                item[key] = value.strftime("%d-%m-%Y")
            elif value is None:
                item[key] = ""
            else:
                item[key] = str(value)

        candidates.append(item)

    wb.close()
    return headers, candidates


@app.route("/")
def index():
    try:
        headers, candidates = get_candidates()
    except Exception as exc:
        flash(str(exc), "error")
        headers, candidates = [], []

    q = request.args.get("q", "").strip().lower()
    status_filter = request.args.get("status", "").strip()

    if q:
        candidates = [
            c for c in candidates
            if q in " ".join(
                str(v) for k, v in c.items()
                if not k.startswith("_")
            ).lower()
        ]

    if status_filter:
        candidates = [
            c for c in candidates
            if c.get("Selection Status", "") == status_filter
        ]

    statuses = [
        "Applied",
        "Under Review",
        "Shortlisted",
        "Interview Scheduled",
        "Selected",
        "Rejected",
    ]

    all_count = len(get_candidates()[1]) if headers else 0

    return render_template(
        "index.html",
        candidates=candidates,
        statuses=statuses,
        total=all_count,
    )


@app.post("/update/<int:row_num>")
def update_candidate(row_num):
    try:
        wb, ws = get_sheet()

        if row_num < 2 or row_num > ws.max_row:
            wb.close()
            flash("Candidate row not found.", "error")
            return redirect(url_for("index"))

        ws.cell(row_num, 7).value = (
            request.form.get("interview_date", "").strip() or None
        )
        ws.cell(row_num, 9).value = (
            request.form.get("status", "Applied").strip()
        )
        ws.cell(row_num, 8).value = (
            request.form.get("ai_explanation", "").strip()
            or "Not generated"
        )

        wb.save(EXCEL_PATH)
        wb.close()
        flash("Candidate details saved to Excel.", "success")

    except Exception as exc:
        flash(f"Could not save changes: {exc}", "error")

    return redirect(url_for("index"))


@app.post("/explain/<int:row_num>")
def explain_project(row_num):
    api_key = os.getenv("OPENAI_API_KEY")

    if not api_key:
        return jsonify({
            "error": (
                "OpenAI API key is missing. "
                "Restart VS Code and check your environment variable."
            )
        }), 400

    try:
        _, candidates = get_candidates()

        candidate = next(
            (c for c in candidates if c["_row"] == row_num),
            None,
        )

        if not candidate:
            return jsonify({"error": "Candidate not found."}), 404

        title = candidate.get("Project Title", "") or "Untitled project"
        github_link = candidate.get("GitHub Link", "").strip()

        if not github_link:
            return jsonify({
                "error": "Please add a public GitHub repository link first."
            }), 400

        # Validate the public GitHub repository URL.
        parsed = urlparse(github_link)
        parts = parsed.path.strip("/").split("/")

        if (
            parsed.scheme != "https"
            or parsed.hostname != "github.com"
            or len(parts) < 2
            or not re.fullmatch(r"[A-Za-z0-9-]+", parts[0])
            or not re.fullmatch(r"[A-Za-z0-9_.-]+", parts[1])
        ):
            return jsonify({
                "error": (
                    "Invalid GitHub link. Use a link like "
                    "https://github.com/owner/repository"
                )
            }), 400

        owner, repo = parts[0], parts[1]

        # Retrieve the repository README from GitHub.
        readme_url = (
            f"https://api.github.com/repos/{owner}/{repo}/readme"
        )

        github_response = requests.get(
            readme_url,
            headers={
                "Accept": "application/vnd.github+json",
                "User-Agent": "Internship-Candidate-Dashboard",
            },
            timeout=15,
        )

        if github_response.status_code == 404:
            return jsonify({
                "error": (
                    "README not found, or the repository is private "
                    "or inaccessible. Check the GitHub link."
                )
            }), 400

        if github_response.status_code == 403:
            return jsonify({
                "error": (
                    "GitHub access was limited. Please try again later."
                )
            }), 502

        github_response.raise_for_status()
        readme_data = github_response.json()

        if readme_data.get("encoding") != "base64":
            return jsonify({
                "error": "The repository README format is unsupported."
            }), 400

        readme_text = base64.b64decode(
            readme_data["content"]
        ).decode("utf-8", errors="replace")

        if not readme_text.strip():
            return jsonify({
                "error": "The repository README is empty."
            }), 400

        # Limit the README text sent to the AI.
        readme_text = readme_text[:12000]

        # Generate the AI explanation.
        client = OpenAI(api_key=api_key)

        ai_response = client.responses.create(
            model="gpt-4.1-mini",
            instructions=(
                "You are a technical project reviewer evaluating "
                "internship candidates. Explain the project in simple "
                "English using these headings: Project Purpose, "
                "Technologies, Main Features, Working Process, and "
                "Limitations or Missing Information. Only state facts "
                "supported by the provided README. Do not claim to have "
                "inspected source code. Do not make hiring decisions "
                "based on personal characteristics. Keep the response "
                "clear, concise, and useful for a human reviewer."
            ),
            input=(
                f"Project title: {title}\n"
                f"GitHub repository: {github_link}\n\n"
                f"README content:\n{readme_text}"
            ),
        )

        explanation = ai_response.output_text.strip()

        if not explanation:
            return jsonify({
                "error": "The AI returned an empty explanation."
            }), 502

        # Save the explanation in Excel column H.
        wb, ws = get_sheet()

        try:
            if row_num < 2 or row_num > ws.max_row:
                return jsonify({"error": "Candidate row not found."}), 404

            ws.cell(row_num, 8).value = explanation
            wb.save(EXCEL_PATH)
        finally:
            wb.close()

        return jsonify({"explanation": explanation})

    except requests.RequestException:
        app.logger.exception("GitHub README request failed")
        return jsonify({
            "error": (
                "Could not access GitHub. Check your internet connection "
                "and whether the repository is public."
            )
        }), 502

    except Exception:
        app.logger.exception("AI project explanation failed")
        return jsonify({
            "error": (
                "AI request failed. Check the VS Code terminal for details. "
                "Verify your API key, API billing, model access, and "
                "internet connection."
            )
        }), 500


if __name__ == "__main__":
    app.run(debug=True)
